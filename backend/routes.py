"""
Looploop API 路由 — v0.45 事件总线架构

REST API + SSE 事件推送。v0.45 关键变更：
- NOTIFICATION_BUS: asyncio.Queue 替代 SQLite 轮询
- _run_pipeline_v045: 单次 async for 驱动 Pipeline，结果就地持久化
"""

import asyncio
import io
import json
import re
import uuid
import zipfile
from typing import Optional

import yaml as _yaml
from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse, Response
from pydantic import BaseModel, Field

from backend.config import settings
from backend.llm_client import LLMClient
from backend.engine.pipeline import Pipeline, PipelineStage
from backend.safe_parser import SafeParser
from backend.database import (
    create_task,
    get_task,
    update_task,
    list_tasks,
    delete_task,
    save_screenplay,
)

router = APIRouter(prefix="/api")

# ═══════════════════════════════════════════════════════════
# v0.45: 内存事件总线 — SSE 响应式推送，零 DB 轮询
# ═══════════════════════════════════════════════════════════
NOTIFICATION_BUS: dict[str, asyncio.Queue] = {}

# 懒加载单例
_llm: Optional[LLMClient] = None
_pipeline: Optional[Pipeline] = None
_parser: Optional[SafeParser] = None


def get_llm() -> LLMClient:
    global _llm
    if _llm is None:
        _llm = LLMClient()
    return _llm


def get_pipeline() -> Pipeline:
    global _pipeline
    if _pipeline is None:
        _pipeline = Pipeline(get_llm())
    return _pipeline


def get_parser() -> SafeParser:
    global _parser
    if _parser is None:
        _parser = SafeParser(get_llm())
    return _parser


# ============================================================
# 请求/响应模型
# ============================================================

class ConvertRequest(BaseModel):
    text: str = Field(..., min_length=100)
    title: Optional[str] = None
    author: Optional[str] = None
    genre: Optional[str] = None


class RegenerateRequest(BaseModel):
    user_feedback: str = ""
    current_screenplay: dict


class ScreenplayUpdate(BaseModel):
    screenplay: dict


# ============================================================
# 路由
# ============================================================

@router.get("/health")
async def health_check():
    return {"status": "ok", "app": settings.app_name,
            "version": "0.45.0", "llm_model": settings.llm_model}


@router.post("/convert")
async def convert_novel(request: ConvertRequest):
    """提交小说文本，启用事件总线并后台运行 Pipeline"""
    task_id = uuid.uuid4().hex[:12]
    title = request.title or "Untitled"

    # 初始化数据库任务 + 事件总线队列
    await create_task(task_id, title, original_text=_normalize_original(request.text))
    NOTIFICATION_BUS[task_id] = asyncio.Queue()

    # 后台异步 Pipeline (单次调用，进度推送入队 + 结果持久化)
    asyncio.create_task(_run_pipeline_v045(task_id, request))

    return {"task_id": task_id}


@router.get("/convert/{task_id}/stream")
async def stream_progress(task_id: str):
    """SSE — 响应式被动等待，由事件总线唤醒，零磁盘 I/O"""
    queue = NOTIFICATION_BUS.get(task_id)
    if queue is None:
        raise HTTPException(status_code=404, detail="任务不存在")

    async def event_generator():
        try:
            while True:
                try:
                    # 15s 心跳：Pipeline 长时间无推送时发 ping 维持链路
                    data = await asyncio.wait_for(queue.get(), timeout=15.0)
                    yield f"data: {json.dumps(data, ensure_ascii=False)}\n\n"
                    if data.get("stage") == "done" or data.get("stage") == "error":
                        break
                except asyncio.TimeoutError:
                    yield ": ping\n\n"  # SSE 注释心跳，保活
        finally:
            NOTIFICATION_BUS.pop(task_id, None)

    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )


@router.get("/screenplay/{task_id}")
async def get_screenplay(task_id: str):
    """获取最终剧本 JSON"""
    task = await get_task(task_id)
    if not task:
        raise HTTPException(status_code=404, detail="任务不存在")
    if task["status"] == "error":
        raise HTTPException(status_code=500, detail=task.get("error", "转换失败"))
    if task["status"] != "done":
        return {"status": task["status"], "progress": task["progress"]}
    return task.get("result", {})


@router.get("/screenplay/{task_id}/export/yaml")
async def export_yaml(task_id: str):
    task = await get_task(task_id)
    if not task or not task.get("result"):
        raise HTTPException(status_code=404, detail="剧本不存在")
    yaml_str = _yaml.dump(task["result"], allow_unicode=True, sort_keys=False,
                          default_flow_style=False, indent=2)
    return Response(content=yaml_str, media_type="application/x-yaml",
                    headers={"Content-Disposition": f"attachment; filename=screenplay_{task_id}.yaml"})


@router.get("/screenplay/{task_id}/export/word")
async def export_word(task_id: str):
    task = await get_task(task_id)
    if not task or not task.get("result"):
        raise HTTPException(status_code=404, detail="剧本不存在")
    from docx import Document
    data = task["result"]
    doc = Document()
    doc.styles["Normal"].font.name = "Microsoft YaHei"
    meta = data.get("meta", {})
    doc.add_heading(meta.get("title", "剧本"), 0)
    doc.add_paragraph(f"原著：{meta.get('source','')}  作者：{meta.get('author','')}")
    doc.add_paragraph(f"类型：{meta.get('genre','')}  总时长：{meta.get('total_duration_estimate','')}")
    doc.add_paragraph("")
    doc.add_heading("角色表", 1)
    table = doc.add_table(rows=1, cols=4)
    for i, text in enumerate(["角色", "类型", "性别", "描述"]):
        table.rows[0].cells[i].text = text
    for char in data.get("characters", []):
        row = table.add_row()
        row.cells[0].text = char.get("canonical_name", "")
        row.cells[1].text = char.get("role_type", "")
        row.cells[2].text = char.get("gender", "")
        row.cells[3].text = char.get("description", "")
    doc.add_paragraph("")
    for scene in data.get("scenes", []):
        doc.add_heading(f"场景 {scene.get('scene_number','')} · {scene.get('location','')}", level=2)
        doc.add_paragraph(f"时间：{scene.get('time_of_day','')}  氛围：{scene.get('atmosphere','')}")
        for shot in scene.get("shots", []):
            p = doc.add_paragraph()
            p.add_run(f"镜头 {shot.get('shot_number','')}  {shot.get('shot_type','')} · {shot.get('camera_movement','')}").bold = True
            dia = shot.get("dialogue")
            if dia: doc.add_paragraph(f"  {dia.get('character','')}：「{dia.get('line','')}」")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return Response(content=buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    headers={"Content-Disposition": f"attachment; filename=screenplay_{task_id}.docx"})


@router.get("/screenplay/{task_id}/export/manifest")
async def export_manifest(task_id: str):
    task = await get_task(task_id)
    if not task or not task.get("result"):
        raise HTTPException(status_code=404, detail="剧本不存在")
    data = task["result"]
    manifest = {"meta": {"title": data.get("meta", {}).get("title", ""), "version": "v0.4"}, "characters": {}, "scenes": {}}
    for c in data.get("characters", []):
        mid = c.get("material_id", c.get("character_id", "?"))
        manifest["characters"][mid] = {"name": c.get("canonical_name", ""), "aliases": c.get("aliases", []),
                                        "role": c.get("role_type", ""), "gender": c.get("gender", ""),
                                        "age": c.get("age_range", ""), "description": c.get("description", "")}
    for s in data.get("scenes", []):
        mid = s.get("material_id", s.get("scene_id", "?"))
        manifest["scenes"][mid] = {"location": s.get("location", ""), "time_of_day": s.get("time_of_day", ""),
                                    "atmosphere": s.get("atmosphere", ""), "shot_count": len(s.get("shots", []))}
    yaml_str = _yaml.dump(manifest, allow_unicode=True, sort_keys=False, default_flow_style=False, indent=2)
    return Response(content=yaml_str, media_type="application/x-yaml",
                    headers={"Content-Disposition": f"attachment; filename=manifest_{task_id}.yaml"})


@router.get("/screenplay/{task_id}/export/zip")
async def export_zip(task_id: str):
    task = await get_task(task_id)
    if not task or not task.get("result"):
        raise HTTPException(status_code=404, detail="剧本不存在")
    data = task["result"]
    title = data.get("meta", {}).get("title", "screenplay")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(f"{title}_剧本.yaml", _yaml.dump(data, allow_unicode=True, sort_keys=False,
                                                      default_flow_style=False, indent=2).encode("utf-8"))
        manifest = {"meta": {"title": title, "version": "v0.4"}, "characters": {}, "scenes": {}}
        for c in data.get("characters", []):
            mid = c.get("material_id", c.get("character_id", "?"))
            manifest["characters"][mid] = {"name": c.get("canonical_name", ""), "aliases": c.get("aliases", []),
                                            "role": c.get("role_type", ""), "gender": c.get("gender", ""),
                                            "age": c.get("age_range", ""), "description": c.get("description", "")}
        for s in data.get("scenes", []):
            mid = s.get("material_id", s.get("scene_id", "?"))
            manifest["scenes"][mid] = {"location": s.get("location", ""), "atmosphere": s.get("atmosphere", ""),
                                        "shot_count": len(s.get("shots", []))}
        zf.writestr(f"{title}_素材清单.yaml", _yaml.dump(manifest, allow_unicode=True, sort_keys=False,
                                                        default_flow_style=False, indent=2).encode("utf-8"))
        from docx import Document
        doc = Document()
        doc.add_heading(title, 0)
        for scene in data.get("scenes", []):
            doc.add_heading(f"S{scene.get('scene_number','')} · {scene.get('location','')}", level=2)
            for shot in scene.get("shots", []):
                p = doc.add_paragraph()
                p.add_run(f"Shot {shot.get('shot_number','')}  {shot.get('shot_type','')} · {shot.get('camera_movement','')}").bold = True
                dia = shot.get("dialogue")
                if dia: doc.add_paragraph(f"  {dia.get('character','')}: 「{dia.get('line','')}」")
        word_buf = io.BytesIO()
        doc.save(word_buf)
        word_buf.seek(0)
        zf.writestr(f"{title}_剧本.docx", word_buf.read())
    buf.seek(0)
    return Response(content=buf.getvalue(), media_type="application/zip",
                    headers={"Content-Disposition": f"attachment; filename={title}_project.zip"})


@router.post("/shots/{scene_id}/{shot_id}/regenerate")
async def regenerate_shot(scene_id: str, shot_id: str, request: RegenerateRequest):
    screenplay = request.current_screenplay
    target_shot, target_scene = None, None
    for scene in screenplay.get("scenes", []):
        for shot in scene.get("shots", []):
            if shot.get("shot_id") == shot_id:
                target_shot, target_scene = shot, scene
                break
    if not target_shot:
        raise HTTPException(status_code=404, detail=f"Shot {shot_id} 不存在")
    context_shots = []
    if target_scene:
        all_shots = target_scene.get("shots", [])
        target_idx = next((i for i, s in enumerate(all_shots) if s.get("shot_id") == shot_id), -1)
        if target_idx >= 0:
            context_shots = all_shots[max(0, target_idx - 2):min(len(all_shots), target_idx + 3)]
    prompt = f"""请重新生成以下镜头。

原始镜头数据：
```json
{json.dumps(target_shot, ensure_ascii=False, indent=2)}
```

用户反馈：{request.user_feedback or '请自由发挥，生成更精彩的版本'}

上下文镜头（仅供参考，不要修改）：
```json
{json.dumps(context_shots, ensure_ascii=False, indent=2)}
```

请只输出这个镜头的合法 JSON。"""
    parser = get_parser()
    try:
        new_shot = await parser.validate_single_shot(
            await get_llm().generate_json(prompt, temperature=0.5)
        )
        return {"shot": new_shot.model_dump()}
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"镜头重刷失败: {e}")


# ============================================================
# 历史管理 + 编辑保存
# ============================================================

@router.get("/tasks")
async def get_tasks():
    return await list_tasks()


@router.patch("/tasks/{task_id}/title")
async def rename_task(task_id: str, body: dict = {}):
    new_title = body.get("title", "").strip()
    if not new_title:
        raise HTTPException(status_code=422, detail="Title cannot be empty")
    task = await get_task(task_id)
    if not task:
        raise HTTPException(status_code=404, detail="任务不存在")
    await update_task(task_id, title=new_title)
    return {"ok": True, "title": new_title}


@router.delete("/tasks/{task_id}")
async def remove_task(task_id: str):
    task = await get_task(task_id)
    if not task:
        raise HTTPException(status_code=404, detail="任务不存在")
    await delete_task(task_id)
    return {"ok": True}


@router.patch("/screenplay/{task_id}")
async def update_screenplay(task_id: str, body: ScreenplayUpdate):
    task = await get_task(task_id)
    if not task:
        raise HTTPException(status_code=404, detail="任务不存在")
    await save_screenplay(task_id, body.screenplay)
    return {"ok": True}


# ============================================================
# v0.45: 事件总线驱动 Pipeline
# ============================================================

async def _run_pipeline_v045(task_id: str, request: ConvertRequest):
    """
    单次 async for 驱动 Pipeline。
    每个阶段：写入数据库持久化 + 推送事件总线 queue。
    最终阶段：Pipeline 产出 Screenplay，就地持久化到 SQLite。
    """
    queue = NOTIFICATION_BUS.get(task_id)
    try:
        pipeline = get_pipeline()
        meta_overrides = {}
        if request.title:
            meta_overrides["title"] = request.title
            meta_overrides["source"] = request.title
        if request.author:
            meta_overrides["author"] = request.author
        if request.genre:
            meta_overrides["genre"] = request.genre

        await update_task(task_id, status="running")

        async for stage in pipeline.run_pipeline(request.text, meta_overrides):
            # 1. 数据持久化到 SQLite
            event_data = {"stage": stage.stage, "message": stage.message,
                          "progress": stage.progress, "data": stage.data}
            if stage.stage == "result":
                # 最终阶段 — 持久化 Screenplay
                sp: Screenplay = stage.data["screenplay"]
                result_dict = sp.model_dump()

                # 回填标题
                final_title = (
                    sp.meta.title if sp.meta.title and sp.meta.title != "Untitled"
                    else request.title or request.text.strip().split("\n")[0][:60]
                )
                await update_task(task_id, title=final_title, result=result_dict,
                                  progress=1.0, stage="done", status="done",
                                  message=stage.message)

                event_data["stage"] = "done"
                event_data["data"] = {"scenes": sp.meta.total_scenes,
                                      "shots": sp.meta.total_shots,
                                      "characters": len(sp.characters),
                                      "duration": sp.meta.total_duration_estimate}
            elif stage.stage == "character_update":
                await update_task(task_id, status="running", stage=stage.stage,
                                  message=stage.message, progress=stage.progress)
                event_data["logs"] = stage.logs
            else:
                await update_task(task_id, status="running", stage=stage.stage,
                                  message=stage.message, progress=stage.progress)

            # 2. 推入事件总线唤醒 SSE 消费者
            await queue.put(event_data)

    except Exception as e:
        if queue:
            await queue.put({"stage": "error", "message": str(e), "progress": 0})
        await update_task(task_id, status="error", error=str(e),
                          message=f"转换失败: {e}")
    finally:
        # 不在这里 pop queue — 由 SSE 消费者 cleanup
        pass


def _normalize_original(text: str) -> str:
    """存储前将内联章节物理拆分为独立段落"""
    t = text.replace("\r\n", "\n").replace("\r", "\n")
    t = re.sub(r"(?<=[。！？!?\n])\s*(第[一二三四五六七八九十百千\d]+[章节回])", r"\n\n\1", t)
    t = re.sub(r"(?<=[.?!\n])\s*(Chapter\s+\d+)", r"\n\n\1", t, flags=re.IGNORECASE)
    return re.sub(r"\n{3,}", "\n\n", t).strip()
